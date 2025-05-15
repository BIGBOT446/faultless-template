from docx import Document
from tkinter import Tk, filedialog


def upload_and_read_docx():
    root = Tk()
    root.withdraw()
    file_path = filedialog.askopenfilename(
        title="Select a Word (.docx) file",
        filetypes=[("Word files", "*.docx")]
    )
    if not file_path:
        print("No file selected.")
        return ""
    try:
        doc = Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        return "\n".join(full_text)
    except Exception as e:
        print(f"Error reading file: {e}")
        return ""


def upload_docx():
    root = Tk()
    root.withdraw()
    file_path = filedialog.askopenfilename(
        title="Select a Word (.docx) file",
        filetypes=[("Word files", "*.docx")]
    )
    if not file_path:
        print("No file selected.")
        return ""
    return str(file_path)


def read_docx(file_path):
    try:
        doc = Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        print("\n".join(full_text))
        return "\n".join(full_text)
    except Exception as e:
        print(f"Error reading file: {e}")
        return ""