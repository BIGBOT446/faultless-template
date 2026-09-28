r"""Extract reviewable text from a .docx file.

``prompt.py`` used ``"\\n".join(p.text for p in Document(report).paragraphs)``.
``Document.paragraphs`` skips table cells entirely, so on a typical engineering
report roughly a quarter of the content never reached the model. Units, bearing
capacities and design parameters live in tables, which is exactly what the
"Units and Symbols Consistency" rule is meant to check.

This module walks the document in order and renders tables as Markdown pipe
tables, so the model sees cell values in their row and column context. The
annotation step needs no change: ``marker/rules.py`` locates text with
Spire's ``FindAllString``, which already searches table cells.
"""

from pathlib import Path

from docx import Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

MAX_CELL_CHARS = 300
DEFAULT_CHUNK_CHARS = 6000
SECTION_STYLES = ("Heading 1", "Title")


def _cell_text(cell: _Cell) -> str:
    """Flatten one cell to a single line, escaping the Markdown delimiter."""
    text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
    text = text.replace("|", "\\|").replace("\n", " ")
    if len(text) > MAX_CELL_CHARS:
        text = text[:MAX_CELL_CHARS] + "..."
    return text


def render_table(table: Table) -> str:
    """Render a table as a Markdown pipe table.

    Merged cells repeat their text across the span, which is what python-docx
    reports; duplicates are kept so column alignment stays intact.
    """
    rows = []
    for row in table.rows:
        rows.append([_cell_text(c) for c in row.cells])
    if not rows:
        return ""

    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]

    header, *body = rows
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


def extract_text(report: str | Path, *, include_tables: bool = True) -> str:
    """Return the document text in reading order.

    Args:
        report: Path to the .docx file.
        include_tables: When False, reproduces the old paragraph-only behaviour,
            which is useful for measuring what the tables add.

    Returns:
        The document body as plain text, with tables as Markdown pipe tables.
    """
    document = Document(str(report))
    parts = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            parts.append(block.text)
        elif isinstance(block, Table) and include_tables:
            rendered = render_table(block)
            if rendered:
                parts.append(rendered)
    return "\n".join(parts)


def _blocks(document: Document, *, include_tables: bool) -> list[tuple[str, str]]:
    """Yield (kind, text) pairs in reading order; kind is "heading", "para" or "table"."""
    out: list[tuple[str, str]] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            kind = "heading" if block.style.name in SECTION_STYLES else "para"
            out.append((kind, block.text))
        elif isinstance(block, Table) and include_tables:
            rendered = render_table(block)
            if rendered:
                out.append(("table", rendered))
    return out


def extract_chunks(
    report: str | Path,
    *,
    include_tables: bool = True,
    max_chars: int = DEFAULT_CHUNK_CHARS,
) -> list[str]:
    """Split the report into section-sized chunks for separate review calls.

    Sections start at a top-level heading. A section longer than ``max_chars`` is
    split further on block boundaries, and every part is prefixed with its section
    heading so the model keeps the context of where the text came from. Tables are
    never split, so a single large table may exceed ``max_chars``.

    Args:
        report: Path to the .docx file.
        include_tables: Pass False to reproduce paragraph-only extraction.
        max_chars: Soft upper bound on the size of one chunk.

    Returns:
        The chunks in document order. Never empty for a non-empty document.
    """
    blocks = _blocks(Document(str(report)), include_tables=include_tables)

    chunks: list[str] = []
    heading = "Front matter"
    current: list[str] = []
    size = 0

    def flush() -> None:
        nonlocal current, size
        if current:
            body = "\n".join(current).strip()
            if body:
                part = len([c for c in chunks if c.startswith(f"[Section: {heading}")]) + 1
                suffix = "" if part == 1 else f" (part {part})"
                label = f"[Section: {heading}{suffix}]"
                chunks.append(f"{label}\n{body}")
        current = []
        size = 0

    for kind, text in blocks:
        if kind == "heading" and text.strip():
            flush()
            heading = text.strip()
            continue
        if not text.strip():
            continue
        if size and size + len(text) > max_chars:
            flush()
        current.append(text)
        size += len(text)
    flush()

    return chunks or [extract_text(report, include_tables=include_tables)]
